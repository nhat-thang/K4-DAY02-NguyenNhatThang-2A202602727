"""experiments.py - các hàm dùng chung cho notebook: phương pháp suy luận (Bước 3), dự đoán chung kết
(Bước 4) và ghi results.xlsx (Bước 5). Mọi lựa chọn dựa trên VAL; test chỉ chạy trong final_predict.
"""
from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd
import torch
from torchvision import transforms as T

import dataset as D
import inference as I
import train
from eval import compute_metrics, save_predictions

TRANSFORMER_PREFIXES = ("vit", "deit", "swin", "beit", "eva")


def is_transformer(backbone: str) -> bool:
    return backbone.startswith(TRANSFORMER_PREFIXES)


@dataclass
class Method:
    exp_id: str
    name: str
    k: int
    input: str                      # "crop" | "full" | "res<r>"
    views: Callable                 # x -> list[batch]
    space: str = "prob"
    cnn_only: bool = False


def methods_for(img_size: int = 224) -> dict[str, Method]:
    """Các phương pháp suy luận một model. I00 là mốc 1-view."""
    m = [
        Method("I00", "1 view (Resize 256 + CenterCrop 224)", 1, "crop", lambda x: [x]),
        Method("I01", "TTA lật ngang, gộp xác suất", 2, "crop", I.views_flip2, "prob"),
        Method("I01L", "TTA lật ngang, gộp logit", 2, "crop", I.views_flip2, "logit"),
        Method("I02", "TTA 5-crop 224 từ ảnh 256, gộp xác suất", 5, "full",
               lambda x: I.views_multicrop(x, img_size), "prob"),
        Method("I02L", "TTA 5-crop 224 từ ảnh 256, gộp logit", 5, "full",
               lambda x: I.views_multicrop(x, img_size), "logit"),
        Method("I02F", "TTA 5-crop + lật (K=10), gộp xác suất", 10, "full",
               lambda x: I.views_multicrop(x, img_size, flip=True), "prob"),
        Method("I02S", "TTA 3 tỉ lệ 224/256/288 từ ảnh 256, gộp xác suất", 3, "full",
               lambda x: I.views_multiscale(x, [224, 256, 288]), "prob", cnn_only=True),
    ]
    for r in (256, 288, 320):
        m.append(Method(f"I04_{r}", f"Độ phân giải kiểm tra {r} (Resize {round(r / 0.875)} + CenterCrop {r})",
                        1, f"res{r}", lambda x: [x], cnn_only=True))
    return {x.exp_id: x for x in m}


def eval_transform(kind: str, img_size: int = 224):
    if kind == "crop":
        return D.build_transforms(False, img_size)
    if kind == "full":   # ảnh gốc 256x256, không cắt
        return T.Compose([T.Resize((256, 256)), T.ToTensor(), T.Normalize(D.IMAGENET_MEAN, D.IMAGENET_STD)])
    if kind.startswith("res"):
        return D.build_transforms(False, int(kind[3:]))
    raise ValueError(kind)


class LoaderCache:
    """Giữ dataset đã nạp bytes vào RAM, chỉ đổi transform -> không đọc đĩa lại cho mỗi phương pháp."""

    def __init__(self, df, images_dir, batch_size=128, num_workers=2):
        self.ds = D.DeepWeedsDataset(df, images_dir, None, preload=True)
        self.batch_size, self.num_workers = batch_size, num_workers

    def loader(self, kind: str, img_size: int = 224):
        self.ds.transform = eval_transform(kind, img_size)
        return torch.utils.data.DataLoader(self.ds, batch_size=self.batch_size, shuffle=False,
                                           num_workers=self.num_workers, pin_memory=True)


def run_method(model, method: Method, cache: LoaderCache, device, img_size: int = 224):
    """Trả về (filenames, y, list logit theo view)."""
    return I.predict_views(model, cache.loader(method.input, img_size), device, method.views)


def probs_of(view_logits, space: str) -> np.ndarray:
    return I.aggregate_views(view_logits, space)


def pseudo_logits(probs: np.ndarray) -> np.ndarray:
    """log p của dự đoán đã gộp; temperature scaling trên log p: softmax(log p / T)."""
    return np.log(np.clip(probs, 1e-12, None))


def metrics(y, probs) -> dict:
    return compute_metrics(np.asarray(y), probs.argmax(1), probs)


def ece_crossfit(logits, y, seed: int = 0) -> float:
    """ECE sau temperature scaling ước lượng trung thực trên val: khớp T trên một nửa, đo trên nửa kia (2 chiều)."""
    rng = np.random.default_rng(seed)
    idx = rng.permutation(len(y))
    a, b = idx[: len(y) // 2], idx[len(y) // 2:]
    eces = []
    for fit, ev in ((a, b), (b, a)):
        t = I.fit_temperature(logits[fit], y[fit])
        p = I.apply_temperature(logits[ev], t)
        eces.append(metrics(y[ev], p)["ece"])
    return float(np.mean(eces))


def final_predict(rd: str | Path, out_exp: str, seed: int, method: Method, pred_dir: str | Path,
                  images_dir, labels_dir, calibrate: bool, device=None, num_workers: int = 2) -> dict:
    """Bước 4: nạp checkpoint tốt nhất (chọn bằng val), chạy phương pháp suy luận đã chốt trên VAL và
    TEST (đúng một lần), khớp T trên VAL, ghi:
        <out_exp>_seed<k>_test.csv  (đã temperature scaling nếu calibrate)
        <out_exp>_uncal_seed<k>_test.csv (nếu calibrate)  ·  <out_exp>_seed<k>_val.csv
    """
    device = device or torch.device("cuda" if torch.cuda.is_available() else "cpu")
    model, cfg = train.load_checkpoint(rd, device)
    _, val_df, test_df = D.load_split(labels_dir, cfg.fold)
    pred_dir = Path(pred_dir)
    out = {"exp_id": out_exp, "seed": seed, "method": method.exp_id, "run_dir": str(rd)}

    names_v, yv, lv = run_method(model, method, LoaderCache(val_df, images_dir, num_workers=num_workers), device,
                                 cfg.img_size)
    pv = probs_of(lv, method.space)
    T_ = 1.0
    if calibrate:
        T_ = I.fit_temperature(pseudo_logits(pv), yv)
        pv_cal = I.apply_temperature(pseudo_logits(pv), T_)
    else:
        pv_cal = pv
    save_predictions(pred_dir / f"{out_exp}_seed{seed}_val.csv", names_v, yv, pv_cal)

    names_t, yt, lt = run_method(model, method, LoaderCache(test_df, images_dir, num_workers=num_workers), device,
                                 cfg.img_size)
    pt = probs_of(lt, method.space)
    if calibrate:
        save_predictions(pred_dir / f"{out_exp}_uncal_seed{seed}_test.csv", names_t, yt, pt)
        pt = I.apply_temperature(pseudo_logits(pt), T_)
    save_predictions(pred_dir / f"{out_exp}_seed{seed}_test.csv", names_t, yt, pt)
    out["T"] = T_
    (Path(rd) / f"final_{out_exp}.json").write_text(json.dumps(out, indent=2), encoding="utf-8")
    return out


# ----------------------------------------------------------------------------- results.xlsx
def write_xlsx(sheets: dict[str, pd.DataFrame], path: str | Path, highlight: dict[str, str] | None = None) -> None:
    """Ghi nhiều sheet: freeze hàng tiêu đề, 4 chữ số thập phân, tô dòng tốt nhất theo cột trong `highlight`."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    highlight = highlight or {}
    fill = PatternFill("solid", fgColor="FFF2CC")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for j, col in enumerate(df.columns, start=1):
                width = max(10, min(60, max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2))
                ws.column_dimensions[get_column_letter(j)].width = width
                if pd.api.types.is_float_dtype(df[col]):
                    for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        row[0].number_format = "0.0000"
            col = highlight.get(name)
            if col and col in df.columns and df[col].notna().any():
                best = int(pd.to_numeric(df[col], errors="coerce").idxmax()) + 2
                for cell in ws[best]:
                    cell.fill = fill


def df_to_md(df: pd.DataFrame, floatfmt: str = ".4f") -> str:
    """Bảng markdown không cần thư viện tabulate."""
    def f(v):
        if isinstance(v, float):
            return "" if np.isnan(v) else format(v, floatfmt)
        return str(v)
    head = "| " + " | ".join(map(str, df.columns)) + " |"
    sep = "|" + "|".join("---" for _ in df.columns) + "|"
    rows = ["| " + " | ".join(f(v) for v in r) + " |" for r in df.itertuples(index=False)]
    return "\n".join([head, sep, *rows])
